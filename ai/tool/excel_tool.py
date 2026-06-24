from pydantic import BaseModel,Field
from langchain.tools import tool
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import BarChart, PieChart, LineChart, ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
import os
import time
import re
load_dotenv()

class ExcelParams(BaseModel):
    content: str = Field(..., description="文档内容(含markdown表格)")

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
table_header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
table_header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
even_row_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
content_font = Font(name='微软雅黑', size=11)
center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)


def parse_tables(content: str):
    tables = []
    table_pattern = r'(\|[^\n]+\|\n\|[-:\s|]+\|\n(?:\|[^\n]+\|\n?)+)'
    for match in re.finditer(table_pattern, content):
        table_text = match.group(0).strip()
        lines = table_text.split('\n')
        headers = [h.strip() for h in lines[0].split('|')[1:-1]]
        data_rows = []
        for line in lines[2:]:
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if any(cells):
                data_rows.append(cells)
        if headers and data_rows:
            tables.append((headers, data_rows))
    return tables


def _find_numeric_cols(headers, data_rows):
    numeric_cols = []
    for ci in range(1, len(headers) + 1):
        is_num = True
        for dr in data_rows:
            try:
                float(dr[ci-1])
            except (ValueError, TypeError):
                is_num = False
                break
        if is_num:
            numeric_cols.append(ci)
    return numeric_cols


def _is_time_series(first_col_values):
    return any(
        re.search(r'(\d{4}[-/年]\d{1,2}[-/月]|\d{4}年|月$|季度$|Q\d|周$)', str(v))
        for v in first_col_values
    )


def _add_bar_chart(ws, table_start_row, table_end_row, numeric_cols, anchor_row):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "柱状图"
    chart.y_axis.title = "数值"
    chart.width = 20
    chart.height = 13

    cats = Reference(ws, min_col=1, min_row=table_start_row + 1, max_row=table_end_row)
    for nci in numeric_cols:
        data_ref = Reference(ws, min_col=nci, min_row=table_start_row, max_row=table_end_row)
        chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{anchor_row}")


def _add_pie_chart(ws, table_start_row, table_end_row, numeric_cols, anchor_row):
    if len(numeric_cols) == 0:
        return
    pie = PieChart()
    pie.title = "饼图"
    pie.width = 14
    pie.height = 13
    pie_data = Reference(ws, min_col=numeric_cols[0], min_row=table_start_row, max_row=table_end_row)
    pie.add_data(pie_data, titles_from_data=True)
    cats = Reference(ws, min_col=1, min_row=table_start_row + 1, max_row=table_end_row)
    pie.set_categories(cats)
    ws.add_chart(pie, f"K{anchor_row}")


def _add_line_chart(ws, table_start_row, table_end_row, numeric_cols, anchor_row, headers):
    line = LineChart()
    line.title = "趋势折线图"
    line.y_axis.title = "数值"
    line.width = 20
    line.height = 13

    cats = Reference(ws, min_col=1, min_row=table_start_row + 1, max_row=table_end_row)
    for nci in numeric_cols:
        data_ref = Reference(ws, min_col=nci, min_row=table_start_row, max_row=table_end_row)
        line.add_data(data_ref, titles_from_data=True)
    line.set_categories(cats)
    ws.add_chart(line, f"A{anchor_row}")


def _add_scatter_chart(ws, table_start_row, table_end_row, numeric_cols, anchor_row, headers, data_rows):
    if len(numeric_cols) < 2:
        return
    col_avgs = []
    for ci in numeric_cols:
        vals = []
        for dr in data_rows:
            try:
                vals.append(float(dr[ci-1]))
            except (ValueError, TypeError):
                pass
        col_avgs.append((ci, sum(vals) / len(vals) if vals else 0))
    col_avgs.sort(key=lambda x: x[1])
    x_col, y_col = col_avgs[0][0], col_avgs[-1][0]

    scatter = ScatterChart()
    scatter.title = "散点图"
    scatter.x_axis.title = headers[x_col - 1]
    scatter.y_axis.title = headers[y_col - 1]
    scatter.width = 18
    scatter.height = 13

    xvalues = Reference(ws, min_col=x_col, min_row=table_start_row + 1, max_row=table_end_row)
    yvalues = Reference(ws, min_col=y_col, min_row=table_start_row + 1, max_row=table_end_row)
    series = Series(yvalues, xvalues, title="数据点")
    scatter.series.append(series)
    ws.add_chart(scatter, f"A{anchor_row}")


@tool("excel_tool",args_schema=ExcelParams)
def excel_tool(content:str)->str:
    """
      将数据写入Excel，自动解析markdown表格并生成统计图表
    """
    try:
        wb = Workbook()
        tables = parse_tables(content)

        if not tables:
            ws = wb.active
            ws.title = "数据"
            ws['A1'] = "未检测到表格数据"
            ws['A1'].font = content_font

        for ti, (headers, data_rows) in enumerate(tables):
            if ti == 0:
                ws = wb.active
            else:
                ws = wb.create_sheet()
            ws.title = f"数据{ti + 1}" if ti > 0 else "数据表"

            num_cols = len(headers)
            table_start_row = 1

            for ci, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=ci, value=header)
                cell.font = table_header_font
                cell.fill = table_header_fill
                cell.alignment = center_align
                cell.border = thin_border
            ws.row_dimensions[1].height = 28

            for ri, data_row in enumerate(data_rows):
                row_idx = ri + 2
                for ci, value in enumerate(data_row, 1):
                    cell = ws.cell(row=row_idx, column=ci)
                    try:
                        cell.value = float(value)
                    except (ValueError, TypeError):
                        cell.value = value
                    cell.font = content_font
                    cell.alignment = center_align
                    cell.border = thin_border
                    if ri % 2 == 1:
                        cell.fill = even_row_fill
                ws.row_dimensions[row_idx].height = 22

            table_end_row = len(data_rows) + 1

            for ci in range(1, num_cols + 1):
                max_width = len(str(headers[ci-1])) * 3
                for dr in data_rows:
                    w = len(str(dr[ci-1])) * 2
                    if w > max_width:
                        max_width = w
                ws.column_dimensions[get_column_letter(ci)].width = min(max(max_width, 10), 35)

            all_num = _find_numeric_cols(headers, data_rows)
            if not all_num or len(data_rows) == 0:
                continue

            cat_num = [ci for ci in all_num if ci > 1]
            has_text_cat = 1 not in all_num

            first_col_values = [dr[0] for dr in data_rows]
            is_time = _is_time_series(first_col_values)
            is_few_rows = len(data_rows) <= 6

            chart_row = table_end_row + 2

            if is_time:
                _add_line_chart(ws, table_start_row, table_end_row, cat_num or all_num, chart_row, headers)
            elif has_text_cat and is_few_rows and len(cat_num) == 1:
                _add_pie_chart(ws, table_start_row, table_end_row, cat_num, chart_row)
            elif len(all_num) >= 2:
                _add_scatter_chart(ws, table_start_row, table_end_row, all_num, chart_row, headers, data_rows)
            else:
                _add_bar_chart(ws, table_start_row, table_end_row, cat_num or all_num, chart_row)

        file_path = os.getenv("FILE_PATH")
        file_name = time.strftime("%Y%m%d%H%M%S", time.localtime())
        full_path = os.path.join(file_path, f"{file_name}.xlsx")
        wb.save(full_path)
        return f"Excel文件路径:{full_path},下载链接:http://localhost:8080/static/{file_name}.xlsx"

    except Exception as e:
        print("写入Excel文档出现异常",e)
        return "写入Excel文档出现异常"

if __name__ =="__main__":
    test_cases = [
        ("时间序列", """| 月份 | 销售额 | 订单数 |
|------|--------|--------|
| 2023年10月 | 150000 | 320 |
| 2023年11月 | 180000 | 400 |
| 2023年12月 | 220000 | 510 |"""),
        ("占比数据", """| 产品类别 | 销售额 |
|------|--------|
| 电子产品 | 50000 |
| 服装 | 30000 |
| 食品 | 20000 |"""),
        ("多变量", """| 广告投入 | 销售额 | 利润 |
|------|--------|--------|
| 10000 | 50000 | 8000 |
| 20000 | 90000 | 15000 |
| 30000 | 120000 | 20000 |
| 40000 | 140000 | 22000 |"""),
    ]
    for name, data in test_cases:
        rs = excel_tool.invoke({"content": data})
        print(f"{name}: {rs}")
