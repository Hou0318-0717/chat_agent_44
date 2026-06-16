from pydantic import BaseModel,Field
from langchain.tools import tool
from docx import Document
from  dotenv import load_dotenv
from pathlib import Path
import time
import os
load_dotenv()

#工具参数类
class WordParams(BaseModel):
    content:str = Field(...,description="文档内容")

'''
写入word
'''
@tool("word_tool",args_schema=WordParams)
def word_tool(content:str)->str:
    """
    写入word文档
    :param content:  文档内容
    :return:  返回一个下载链接
    """
    try:
        #创建一个文档对象
        docx = Document()
        #写入标题
        docx.add_heading("文档标题",level=1)
        #写入正文
        docx.add_paragraph(content)
        #获取文件路径
        file_path = os.getenv("FILE_PATH")
        #为了保证文件名唯一性
        file_name = time.strftime("%Y%m%d%H%M%S",time.localtime())
        #保存文档
        docx.save(f"{file_path}/{file_name}.docx")
        return f"文档下载链接是:http://localhost:8080/static/{file_name}.docx"
    except Exception as e:
        print("写入word文档出现异常",e)
        return "写入word文档出现异常"

if __name__ == '__main__':
    word_tool.invoke({"content":"这是一个测试内容"})
